import json
import sqlite3
import struct
import zipfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import duckdb
import pytest

from zhishu_inspection.file_sources import read_file_tables
from zhishu_inspection import file_sources


@pytest.mark.parametrize('extension,separator,encoding', [('csv', ',', 'utf-8-sig'), ('tsv', '\t', 'utf-8'), ('txt', '|', 'gb18030'), ('txt', ';', 'utf-8')])
def test_delimited_text_preserves_strings(tmp_path, extension, separator, encoding):
    path = tmp_path / ('订单.' + extension)
    path.write_bytes((separator.join(['编号', '金额']) + '\n' + separator.join(['0001', '12.30']) + '\n').encode(encoding))
    table, = read_file_tables(path)
    assert table.columns == ('编号', '金额')
    assert table.rows == [('0001', '12.30')]
    assert table.origin['format'] == extension


@pytest.mark.parametrize('extension,content', [('json', '[{"id":"0001","amount":12.30}]'), ('json', '{"records":[{"id":"0001","amount":12.30}]}'), ('jsonl', '{"id":"0001","amount":12.30}\n'), ('ndjson', '{"id":"0001","amount":12.30}\n')])
def test_json_readers_keep_decimal_precision(tmp_path, extension, content):
    path = tmp_path / ('data.' + extension)
    path.write_text(content, encoding='utf-8')
    table, = read_file_tables(path)
    assert table.rows == [('0001', Decimal('12.30'))]


@pytest.mark.parametrize('content', ['[{"id":1,"nested":{"x":2}}]', '[{"id":1,"nested":[2]}]', '{"rows":[{"id":1}]}', '[{"id":1},{"other":2}]', '[{"id":1,"id":2}]'])
def test_json_rejects_nested_ambiguous_or_inconsistent_records(tmp_path, content):
    path = tmp_path / 'data.json'
    path.write_text(content, encoding='utf-8')
    with pytest.raises(ValueError):
        read_file_tables(path)


@pytest.mark.parametrize('extension,content', [('csv', b'id,id\n1,2'), ('csv', b'id,\n1,2'), ('csv', b'id,value\n1,2,3'), ('csv', b'id,value\n"broken,2'), ('json', b'[]'), ('json', b'{bad'), ('csv', b''), ('pdf', b'%PDF')])
def test_rejects_invalid_or_unsupported_files(tmp_path, extension, content):
    path = tmp_path / ('data.' + extension)
    path.write_bytes(content)
    with pytest.raises(ValueError):
        read_file_tables(path)


def test_readers_bound_rows_and_bytes_without_truncation(tmp_path, monkeypatch):
    path = tmp_path / 'data.csv'
    path.write_text('id\n1\n2\n3\n', encoding='utf-8')
    monkeypatch.setattr(file_sources, 'MAX_ROWS', 2)
    with pytest.raises(ValueError, match='行数'):
        read_file_tables(path)
    monkeypatch.setattr(file_sources, 'MAX_BYTES', 3)
    with pytest.raises(ValueError, match='大小'):
        read_file_tables(path)


def test_sqlite_file_lists_user_tables_and_quotes_names(tmp_path):
    path = tmp_path / 'data.sqlite'
    with sqlite3.connect(path) as conn:
        conn.execute('create table "a\""b" (id text, amount text)')
        conn.execute('insert into "a\""b" values (?,?)', ('0001', '12.30'))
        conn.execute('create view ignored as select * from "a\""b"')
    table, = read_file_tables(path)
    assert table.name == 'a"b'
    assert table.rows == [('0001', '12.30')]
    assert table.origin['table'] == 'a"b'


def test_sqlite_upload_omits_virtual_tables_and_fts_shadow_tables(tmp_path):
    path = tmp_path / 'search.sqlite'
    with sqlite3.connect(path) as conn:
        conn.execute('create table orders (id text)')
        conn.execute("insert into orders values ('0001')")
        conn.execute('create virtual table search_index using fts5(content)')
        conn.execute("insert into search_index values ('searchable text')")
    tables = read_file_tables(path)
    assert [table.name for table in tables] == ['orders']
    assert tables[0].rows == [('0001',)]


def test_sqlite_reader_observes_committed_wal_only_changes(tmp_path):
    path = tmp_path / 'live.sqlite'
    writer = sqlite3.connect(path)
    try:
        writer.execute('PRAGMA journal_mode=WAL')
        writer.execute('create table orders (id text, amount text)')
        writer.execute("insert into orders values ('0001','100.00')")
        writer.commit()
        writer.execute('PRAGMA wal_checkpoint(TRUNCATE)')
        original_main_file = path.read_bytes()
        writer.execute("update orders set amount='150.00'")
        writer.commit()
        assert path.read_bytes() == original_main_file
        table, = read_file_tables(path)
        assert table.rows == [('0001', '150.00')]
        # The file reader never checkpoints or writes back into the database.
        assert path.read_bytes() == original_main_file
    finally:
        writer.close()


def test_gb18030_csv_cannot_expand_past_utf8_table_budget(tmp_path):
    path = tmp_path / 'expanded.csv'
    # Each Chinese character occupies two GB18030 bytes but three UTF-8 bytes.
    path.write_bytes(('id\n' + ('中' * 900 + '\n') * 8000).encode('gb18030'))
    assert path.stat().st_size < file_sources.MAX_BYTES
    with pytest.raises(ValueError, match='20 MB|字节|大小'):
        read_file_tables(path)


def test_duckdb_and_parquet_preserve_native_decimal_types(tmp_path):
    db = tmp_path / 'data.duckdb'
    parquet = tmp_path / 'data.parquet'
    with duckdb.connect(str(db)) as conn:
        conn.execute("create table orders as select '0001' id, 12.30::decimal(10,2) amount")
        conn.execute('copy orders to ? (format parquet)', [str(parquet)])
    for path in (db, parquet):
        table, = read_file_tables(path)
        assert table.columns == ('id', 'amount')
        assert table.rows == [('0001', Decimal('12.30'))]


def test_xlsx_reads_each_nonempty_sheet_and_dates(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / 'data.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '订单'
    sheet.append(['id', 'date'])
    sheet.append(['0001', datetime(2026, 1, 1)])
    workbook.create_sheet('empty')
    workbook.save(path)
    table, = read_file_tables(path)
    assert table.name == '订单'
    assert table.rows[0][0] == '0001'
    assert table.rows[0][1].year == 2026


def test_xls_reads_real_legacy_workbook(tmp_path):
    import xlwt
    path = tmp_path / 'data.xls'
    book = xlwt.Workbook()
    sheet = book.add_sheet('orders')
    for row_index, row in enumerate([['id', 'amount'], ['0001', '12.30']]):
        for column_index, value in enumerate(row):
            sheet.write(row_index, column_index, value)
    book.save(str(path))
    assert read_file_tables(path)[0].rows == [('0001', '12.30')]


def test_legacy_workbook_rejects_large_offset_before_materializing(tmp_path):
    import xlwt
    path = tmp_path / 'offset.xls'
    book = xlwt.Workbook()
    sheet = book.add_sheet('orders')
    sheet.write(30000, 0, 'id')
    sheet.write(30001, 0, '0001')
    book.save(str(path))
    with pytest.raises(ValueError, match='行数'):
        read_file_tables(path)


def test_ods_reads_real_workbook(tmp_path):
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableRow, TableCell
    from odf.text import P
    path = tmp_path / 'data.ods'
    book = OpenDocumentSpreadsheet()
    sheet = Table(name='orders')
    for values in [['id', 'amount'], ['0001', '12.30']]:
        row = TableRow()
        for value in values:
            cell = TableCell(valuetype='string')
            cell.addElement(P(text=value))
            row.addElement(cell)
        sheet.addElement(row)
    book.spreadsheet.addElement(sheet)
    book.save(str(path))
    assert read_file_tables(path)[0].rows == [('0001', '12.30')]


def test_xlsb_reads_real_upstream_binary_workbook():
    # Calamine upstream regression fixture, containing a single text cell.
    path = Path(__file__).parent / 'fixtures/file_formats/calamine_issue_419.xlsb'
    table, = read_file_tables(path)
    assert table.name == 'Sheet1'
    assert table.columns == ('Hello',)
    assert table.rows == []
    assert table.origin['format'] == 'xlsb'


@pytest.mark.parametrize('field,value', [('dimension_row', 40000), ('dimension_column', 1000), ('actual_row', 40000), ('actual_column', 1000), ('relocated_row', 40000)])
def test_xlsb_coordinates_are_rejected_before_native_reader(tmp_path, monkeypatch, field, value):
    import python_calamine
    original = Path(__file__).parent / 'fixtures/file_formats/calamine_issue_419.xlsb'
    path = tmp_path / 'oversized.xlsb'
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(path, 'w') as target:
        for member in source.infolist():
            data = source.read(member)
            if member.filename == 'xl/worksheets/sheet1.bin':
                changed = bytearray(data)
                # Genuine upstream fixture: BrtWsDim(32), BrtRowHdr(149), BrtCellIsst(176).
                offset = {'dimension_row': 36, 'dimension_column': 44,
                          'actual_row': 149, 'actual_column': 176, 'relocated_row': 149}[field]
                struct.pack_into('<I', changed, offset, value)
                data = bytes(changed)
            if field == 'relocated_row' and member.filename == 'xl/_rels/workbook.bin.rels':
                data = data.replace(b'worksheets/sheet1.bin', b'hidden_sheet.bin')
            name = 'xl/hidden_sheet.bin' if field == 'relocated_row' and member.filename == 'xl/worksheets/sheet1.bin' else member.filename
            target.writestr(name, data)
    called = []
    class NativeReader:
        @staticmethod
        def from_path(path):
            called.append(path)
            raise AssertionError('native parser must not receive oversized coordinates')
    monkeypatch.setattr(python_calamine, 'CalamineWorkbook', NativeReader)
    with pytest.raises(ValueError, match='行数|列数'):
        read_file_tables(path)
    assert called == []


@pytest.mark.parametrize('field', ['dimension_row', 'dimension_column', 'actual_row', 'actual_column', 'old_biff'])
def test_xls_biff8_preflight_blocks_large_coordinates_and_old_versions(tmp_path, monkeypatch, field):
    import olefile
    import python_calamine
    import xlwt
    path = tmp_path / 'oversized.xls'
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet('orders')
    sheet.write(0, 0, 'id')
    sheet.write(1, 0, '0001')
    workbook.save(str(path))
    with olefile.OleFileIO(path, write_mode=True) as archive:
        content = bytearray(archive.openstream('Workbook').read())
        offset = 0
        changed = False
        while offset + 4 <= len(content):
            record_type, size = struct.unpack_from('<HH', content, offset)
            start = offset + 4
            if record_type == 0x200 and field.startswith('dimension_'):
                if field == 'dimension_row':
                    struct.pack_into('<I', content, start + 4, 40001)
                else:
                    struct.pack_into('<H', content, start + 10, 201)
                changed = True
                break
            if record_type == 0xfd and field.startswith('actual_'):
                struct.pack_into('<H', content, start + (0 if field == 'actual_row' else 2),
                                 40000 if field == 'actual_row' else 200)
                changed = True
                break
            if record_type == 0x809 and field == 'old_biff':
                struct.pack_into('<H', content, start, 0x500)
                changed = True
                break
            offset = start + size
        assert changed
        archive.write_stream('Workbook', bytes(content))
    called = []
    class NativeReader:
        @staticmethod
        def from_path(path):
            called.append(path)
            raise AssertionError('preflight must reject before native parsing')
    monkeypatch.setattr(python_calamine, 'CalamineWorkbook', NativeReader)
    with pytest.raises(ValueError, match='行数|列数|BIFF8'):
        read_file_tables(path)
    assert called == []


def test_xlsx_rejects_excessive_archive_expansion(tmp_path, monkeypatch):
    path = tmp_path / 'bomb.xlsx'
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('xl/worksheets/sheet1.xml', ' ' * 10000)
    monkeypatch.setattr(file_sources, 'MAX_ARCHIVE_BYTES', 1000)
    with pytest.raises(ValueError, match='解压'):
        read_file_tables(path)


def test_reader_errors_never_reveal_source_data(tmp_path):
    path = tmp_path / 'secret.csv'
    path.write_text('id,value\nsecret-credential,2,3', encoding='utf-8')
    with pytest.raises(ValueError) as error:
        read_file_tables(path)
    assert 'secret-credential' not in str(error.value)


def test_total_rows_are_bounded_across_database_tables(tmp_path, monkeypatch):
    path = tmp_path / 'data.sqlite'
    with sqlite3.connect(path) as conn:
        for name in ('first', 'second'):
            conn.execute(f'create table {name} (id text)')
            conn.executemany(f'insert into {name} values (?)', [('1',), ('2',)])
    monkeypatch.setattr(file_sources, 'MAX_ROWS', 3)
    with pytest.raises(ValueError, match='合计'):
        read_file_tables(path)


def test_header_only_source_is_previewable(tmp_path):
    path = tmp_path / 'empty.csv'
    path.write_text('id,amount\n', encoding='utf-8')
    table, = read_file_tables(path)
    assert table.columns == ('id', 'amount')
    assert table.rows == []


def test_compressed_parquet_cell_expansion_is_bounded(tmp_path, monkeypatch):
    path = tmp_path / 'large.parquet'
    with duckdb.connect() as conn:
        conn.execute("copy (select repeat('a',10000) as payload) to ? (format parquet)", [str(path)])
    assert path.stat().st_size < 2000
    monkeypatch.setattr(file_sources, 'MAX_BYTES', 2000)
    with pytest.raises(ValueError, match='展开|大小'):
        read_file_tables(path)


@pytest.mark.parametrize('member,content', [
    ('xl/worksheets/sheet1.xml', '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><dimension ref="A1:XFD1048576"/></worksheet>'),
    ('content.xml', '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"><table:table><table:table-row table:number-rows-repeated="1000000000"><table:table-cell/></table:table-row></table:table></office:document-content>'),
])
def test_rejects_synthetic_workbook_dimensions_before_parsing(tmp_path, member, content):
    path = tmp_path / ('data.ods' if member == 'content.xml' else 'data.xlsx')
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr(member, content)
    with pytest.raises(ValueError, match='行数|列数'):
        read_file_tables(path)
